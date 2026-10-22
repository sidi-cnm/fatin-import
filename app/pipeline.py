"""
Les étapes du traitement d'UNE unité, dans l'ordre :

  1. lire le fichier (PDF envoyé tel quel ; Word converti en texte)
  2. demander le découpage à l'IA (avec nouvelles tentatives si le serveur est surchargé)
  3. nettoyer le texte arabe dans tout le JSON
  4. vérifier le FORMAT (Pydantic) -> si raté, 1 seconde chance à l'IA avec les erreurs
  5. vérifier les RÈGLES pédagogiques -> liste de problèmes pour l'écran de relecture

Ce fichier contient aussi les outils partagés avec le découpage automatique d'un livre
entier (outline.py, lesson.py, assemble.py) : découpe des pages, appel à l'IA avec
nouvelles tentatives, lecture du JSON de l'IA.
"""
import io
import json
import re
import time
from pathlib import Path
from typing import Callable, Optional, TypeVar

from pydantic import BaseModel, ValidationError

from .text_cleaning import normalize_tree
from .prompts import build_instructions, build_repair_message
from .providers import AIProvider, QuotaExhaustedError, SourceDocument, TransientProviderError
from .rules import check_draft
from .schemas import ImportDraft, Issue, Meta, ParseResponse, ValidateResponse

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT = "text/plain"
_EXTENSIONS = {".pdf": PDF, ".docx": DOCX, ".txt": TEXT}

Progress = Optional[Callable[[str], None]]
M = TypeVar("M", bound=BaseModel)


class UnsupportedFile(ValueError):
    pass


class PipelineError(RuntimeError):
    """L'IA n'a pas produit un JSON valide, même après la seconde chance."""

    def __init__(self, errors: str, raw_output: str):
        super().__init__(errors)
        self.errors = errors
        self.raw_output = raw_output


# --- 1. Lecture du fichier ---------------------------------------------------------
def read_document(filename: str, data: bytes) -> SourceDocument:
    mime = _EXTENSIONS.get(Path(filename).suffix.lower())
    if mime is None:
        raise UnsupportedFile("Format accepté : .pdf, .docx ou .txt")
    if mime == PDF:
        if not data.startswith(b"%PDF"):
            raise UnsupportedFile("Le fichier n'est pas un PDF valide.")
        return SourceDocument(filename=filename, mime_type=PDF, data=data)
    if mime == DOCX:
        return SourceDocument(filename=filename, mime_type=DOCX, data=data, text=_docx_to_text(data))
    return SourceDocument(filename=filename, mime_type=TEXT, data=data, text=data.decode("utf-8", errors="replace"))


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


# --- Outils PDF : compter, garder certaines pages, découper page par page ------------
MAX_PAGES_WITHOUT_SELECTION = 40
# Poids maximal des pages envoyées une par une (en PDF, sinon en images).
MAX_LABELED_BYTES = 18 * 1024 * 1024


def _open_pdf(data: bytes):
    from pypdf import PdfReader

    if not data.startswith(b"%PDF"):
        raise UnsupportedFile("Le fichier n'est pas un PDF valide.")
    try:
        return PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedFile(f"PDF illisible : {exc}") from exc


def _write_pdf(pages) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for page in pages:
        writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def count_pdf_pages(data: bytes) -> int:
    return len(_open_pdf(data).pages)


def extract_pdf_pages(data: bytes, spec: str) -> bytes:
    """Garde seulement certaines pages d'un PDF. spec : "5-24" ou "5-10,12,15-18" (pages comptées à partir de 1)."""
    reader = _open_pdf(data)
    total = len(reader.pages)
    wanted: list[int] = []
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        match = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not match:
            raise UnsupportedFile(f"Pages invalides : « {part} ». Exemple : 5-24 ou 5-10,12")
        start = int(match.group(1))
        end = int(match.group(2) or start)
        if start < 1 or end < start or end > total:
            raise UnsupportedFile(f"Pages « {part} » hors du document (1 à {total}).")
        wanted.extend(range(start, end + 1))
    if not wanted:
        raise UnsupportedFile("Aucune page sélectionnée.")
    return _write_pdf(reader.pages[n - 1] for n in wanted)


def labeled_pdf_document(filename: str, data: bytes, start: int, end: int) -> SourceDocument:
    """Prépare les pages start..end d'un PDF pour l'IA, chacune avec son numéro dans le livre.

    - `data` : le bloc de pages start..end (un seul PDF) -> utilisé par Mistral (OCR)
    - `pages` : les mêmes pages, une par fichier -> Gemini les reçoit précédées de « === PAGE n === »
      * d'abord en PDF d'une page ;
      * si c'est trop lourd (livre plein d'images), en images JPEG légères (100 dpi) ;
      * si c'est encore trop lourd ou impossible : pas d'étiquettes, l'IA reçoit le bloc
        et la consigne lui dit que la 1re page est la page `start`.
    - `first_page` : numéro de la 1re page dans le livre, pour que les numéros restent ceux du livre
    """
    reader = _open_pdf(data)
    total = len(reader.pages)
    if start < 1 or end < start or end > total:
        raise UnsupportedFile(f"Pages {start}-{end} hors du document (1 à {total}).")
    selected = [reader.pages[n - 1] for n in range(start, end + 1)]
    block = _write_pdf(selected)
    doc = SourceDocument(filename=filename, mime_type=PDF, data=block, first_page=start)

    pages: list[bytes] = []
    size = 0
    for page in selected:
        single = _write_pdf([page])
        size += len(single)
        if size > MAX_LABELED_BYTES:
            break
        pages.append(single)
    else:
        doc.pages = pages
        return doc

    images = render_pages_jpeg(block)
    if images and sum(len(i) for i in images) <= MAX_LABELED_BYTES:
        doc.pages, doc.pages_mime = images, "image/jpeg"
    return doc


def render_pages_jpeg(data: bytes, dpi: int = 100, quality: int = 70) -> Optional[list[bytes]]:
    """Chaque page du PDF en image JPEG légère. None si pypdfium2 n'est pas installé."""
    from .providers.images import render_pdf_pages  # un seul appel à pdfium à la fois

    try:
        return render_pdf_pages(data, dpi, quality)
    except Exception:  # pypdfium2 absent ou PDF illisible : on garde les pages en PDF
        return None


def read_bookmarks(data: bytes, limit: int = 300) -> list[tuple[str, int]]:
    """Les signets du PDF (titre, page), s'il y en a. Gratuit, sans IA : un indice pour le plan."""
    try:
        reader = _open_pdf(data)
        found: list[tuple[str, int]] = []

        def walk(items, depth: int) -> None:
            for item in items:
                if isinstance(item, list):
                    walk(item, depth + 1)
                    continue
                try:
                    found.append(("  " * depth + str(item.title), reader.get_destination_page_number(item) + 1))
                except Exception:
                    continue

        walk(reader.outline, 0)
        return found[:limit]
    except Exception:
        return []


# --- Appel à l'IA : nouvelles tentatives + lecture du JSON ---------------------------
def generate_with_retry(
    provider: AIProvider,
    doc: SourceDocument,
    instructions: str,
    repair: Optional[str],
    attempts: int,
    wait_seconds: float,
    progress: Progress,
) -> str:
    """Appelle l'IA ; si le serveur est surchargé, attend puis réessaie (15 s, 30 s, 60 s...).

    - quota du jour épuisé : QuotaExhaustedError tout de suite (attendre ne sert à rien) ;
    - quota par minute (429) : on attend au moins le délai conseillé par le fournisseur ;
      si le 429 est toujours là après tous les essais, c'est un quota épuisé : QuotaExhaustedError.
    """
    attempts = max(1, attempts)
    for attempt in range(1, attempts + 1):
        try:
            return provider.generate(doc, instructions, repair)
        except TransientProviderError as exc:
            if attempt == attempts:
                if exc.quota:
                    raise QuotaExhaustedError(f"Quota épuisé (429 après {attempts} essais). {exc}") from exc
                raise
            wait = wait_seconds * (2 ** (attempt - 1))
            if exc.retry_after:
                wait = max(wait, exc.retry_after + 1)
            reason = "Quota par minute atteint" if exc.quota else "Serveur d'IA occupé"
            if progress:
                progress(f"{reason} (essai {attempt}/{attempts}). Nouvel essai dans {wait:.0f} s... [{exc}]"[:300])
            time.sleep(wait)
    raise AssertionError("inaccessible")


def parse_json(raw: str):
    text = (raw or "").strip()
    fenced = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fenced:
        text = fenced.group(1)
    try:
        return json.loads(text)
    except ValueError:
        # certains modèles ajoutent une phrase avant ou après le JSON : on garde l'objet { ... }
        start, end = text.find("{"), text.rfind("}")
        if start == -1 or end <= start:
            raise
        return json.loads(text[start:end + 1])


def format_errors(exc: ValidationError, limit: int = 30) -> str:
    lines = []
    for e in exc.errors()[:limit]:
        location = ".".join(str(part) for part in e["loc"])
        lines.append(f"- {location}: {e['msg']}")
    if exc.error_count() > limit:
        lines.append(f"- ... et {exc.error_count() - limit} autres erreurs")
    return "\n".join(lines)


def to_model(raw: str, model: type[M], force: Optional[dict] = None) -> tuple[Optional[M], str]:
    """JSON brut de l'IA -> objet vérifié. Renvoie (objet, "") ou (None, "liste des erreurs").

    Le texte est nettoyé (arabe, français) et les valeurs de `force` sont imposées.
    """
    try:
        data = parse_json(raw)
    except ValueError as exc:
        return None, f"JSON illisible : {exc}"
    if not isinstance(data, dict):
        return None, "La racine du JSON doit être un objet."
    data = normalize_tree(data)
    data.update(force or {})
    try:
        return model.model_validate(data), ""
    except ValidationError as exc:
        return None, format_errors(exc)


def ask_for_model(
    provider: AIProvider,
    doc: SourceDocument,
    instructions: str,
    model: type[M],
    force: Optional[dict] = None,
    retry_attempts: int = 4,
    retry_wait_seconds: float = 15,
    progress: Progress = None,
) -> tuple[M, str, bool]:
    """Demande un objet à l'IA ; si le FORMAT est faux, 1 seconde chance avec les erreurs.

    Renvoie (objet, JSON brut, seconde_chance_utilisée). Lève PipelineError après 2 échecs.
    """

    def ask(repair: Optional[str] = None) -> str:
        return generate_with_retry(provider, doc, instructions, repair, retry_attempts, retry_wait_seconds, progress)

    raw = ask()
    obj, errors = to_model(raw, model, force)
    if not errors:
        return obj, raw, False
    if progress:
        progress("Le JSON de l'IA n'est pas conforme : seconde chance avec la liste des erreurs...")
    raw = ask(build_repair_message(raw, errors))
    obj, errors = to_model(raw, model, force)
    if errors:
        raise PipelineError(errors, raw)
    return obj, raw, True


# --- 2 à 5. Traitement complet d'une unité -------------------------------------------
def run_pipeline(
    doc: SourceDocument,
    provider: AIProvider,
    subject_id: int,
    language: str = "ar",
    stage: str = "",
    level: str = "",
    subject_name: str = "",
    retry_attempts: int = 4,
    retry_wait_seconds: float = 15,
    progress: Progress = None,
) -> ParseResponse:
    started = time.monotonic()
    instructions = build_instructions(subject_id, doc.filename, language, stage, level, subject_name)
    force = {"subject": subject_id, "source_file": doc.filename, "language": language}
    draft, _, repaired = ask_for_model(
        provider, doc, instructions, ImportDraft, force, retry_attempts, retry_wait_seconds, progress
    )
    return ParseResponse(
        draft=draft,
        issues=check_draft(draft),
        meta=Meta(
            provider=provider.name,
            model=provider.model,
            duration_seconds=round(time.monotonic() - started, 2),
            repaired=repaired,
        ),
    )


# --- Vérification d'un brouillon déjà prêt (sans IA) --------------------------------
def validate_draft_data(data: dict) -> ValidateResponse:
    """Vérifie un JSON (format + règles). Utile aux développeurs FATIN et pour les tests."""
    try:
        draft = ImportDraft.model_validate(data)
    except ValidationError as exc:
        issues = [
            Issue(
                code="FORMAT",
                level="error",
                path=".".join(str(p) for p in e["loc"]),
                message=e["msg"],
            )
            for e in exc.errors()
        ]
        return ValidateResponse(valid=False, issues=issues)
    issues = check_draft(draft)
    return ValidateResponse(valid=not any(i.level == "error" for i in issues), issues=issues)
