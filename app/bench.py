"""
Benchmark : comparer les modèles d'IA sur NOTRE tâche (détailler des leçons d'un vrai manuel).

    # 1. lancer les mêmes leçons avec plusieurs modèles (unité 1 par défaut)
    python3 -m app.bench run "$LIVRE" --plan "$PLAN" \\
        --models gemini:gemini-3.8-flash,openai:gpt-6.1-sol,anthropic:claude-sonnet-5-5

    # 2. après la notation à l'aveugle par l'admin (fichier Excel rempli)
    python3 -m app.bench note resultats/bench/revue_aveugle.xlsx

Ce qui est mesuré, pour chaque modèle :
  - QUALITÉ automatique : réponses fausses marquées justes (grave) et bonnes réponses refusées,
    comparées au corrigé de référence (bench/corrige_*.json, vérifié sur le livre) ;
    erreurs de règles (rules.py), leçons ratées, images à ajouter, réponses « VERIFY ».
  - COÛT RÉEL : les tokens renvoyés par l'API × le prix du modèle (app/prices.py).
  - DURÉE par leçon, pannes (503, quota).
  - QUALITÉ vue par l'admin : une feuille Excel où les chunks de tous les modèles sont mélangés,
    SANS le nom du modèle ; l'admin écrit garder / modifier / supprimer.

Chaque leçon faite est enregistrée (resultats/bench/<modèle>/run<n>/) : si un quota s'épuise,
relancer la même commande reprend là où elle s'était arrêtée, et les autres modèles continuent.
"""
import argparse
import json
import random
import re
import sys
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Optional

from pydantic import ValidationError

from .config import get_settings
from .lesson import lesson_fingerprint, run_lesson
from .pipeline import PipelineError
from .prices import PRICES_DATE, cost_of, price_of
from .providers import ProviderError, QuotaExhaustedError, TransientProviderError, Usage, get_provider
from .schemas import LessonResponse, Outline

OUT = Path("resultats") / "bench"
DEFAULT_GOLD = Path("bench") / "corrige_maths_1AF_unite1.json"
DEFAULT_MODELS = (
    "gemini:gemini-3.8-flash,gemini:gemini-3.5-flash-lite,"
    "openai:gpt-6.1-sol,openai:gpt-6-luna,"
    "anthropic:claude-sonnet-5-5,anthropic:claude-haiku-5-5,"
    "mistral:mistral-medium-latest"
)


# --- Lecture du corrigé : des « faits » vrais ou faux, vérifiés sur le livre ----------------
_HARAKAT = re.compile(r"[ً-ْٰـ]")
_PREFIXES = ("وال", "فال", "بال", "كال", "لل", "ال", "و", "ف", "ب", "ل", "ك")
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


def normalize(text: str) -> str:
    text = _HARAKAT.sub("", text or "").translate(_DIGITS)
    for a, b in (("أ", "ا"), ("إ", "ا"), ("آ", "ا"), ("ٱ", "ا"), ("ة", "ه"), ("ى", "ي")):
        text = text.replace(a, b)
    return text


def tokens(text: str) -> list[set[str]]:
    """Chaque mot, avec ses formes sans préfixe (« والكرة » -> {والكره, الكره, كره}...)."""
    out = []
    for word in re.findall(r"[ء-ي]+|\d+|[a-zA-Z]+", normalize(text)):
        forms = {word}
        for prefix in _PREFIXES:
            if word.startswith(prefix) and len(word) - len(prefix) >= 2:
                forms.add(word[len(prefix):])
        out.append(forms)
    return out


def _word_matches(pattern: str, forms: set[str]) -> bool:
    for alt in normalize(pattern).split("|"):
        if alt.endswith("*"):
            if any(f.startswith(alt[:-1]) for f in forms):
                return True
        elif alt in forms:
            return True
    return False


def ordered_match(words: list[str], toks: list[set[str]]) -> Optional[list[int]]:
    """Positions des mots du fait, dans cet ordre, ou None."""
    positions, start = [], 0
    for w in words:
        for i in range(start, len(toks)):
            if _word_matches(w, toks[i]):
                positions.append(i)
                start = i + 1
                break
        else:
            return None
    return positions


def any_order_match(words: list[str], toks: list[set[str]]) -> Optional[list[int]]:
    """Positions des mots du fait, dans n'importe quel ordre (chaque mot sur un mot différent), ou None."""
    positions: list[int] = []
    for w in words:
        found = next((i for i, t in enumerate(toks) if i not in positions and _word_matches(w, t)), None)
        if found is None:
            return None
        positions.append(found)
    return positions


def match_fact(words: list[str], toks: list[set[str]]) -> Optional[list[int]]:
    """CORRECTION 2 — l'ordre des mots.

    - Fait de 2 mots (un nombre et une chose : « 2 + pommes », « banane + main droite ») :
      l'ordre ne change pas le sens. En arabe on écrit « تفاحتين اثنتين » (pommes deux) :
      on accepte les deux ordres.
    - Fait de 3 mots ou plus (« voiture + droite + ballon ») : l'ordre CHANGE le sens
      (« le ballon à droite de la voiture » dit le contraire) : on garde l'ordre.
    """
    if len(words) <= 2:
        return any_order_match(words, toks)
    return ordered_match(words, toks)


# CORRECTION 1 — les questions avec « ne ... pas » (« quel groupe N'a PAS 4 éléments ? »)
_NEGATIONS = {"لا", "ليس", "ليست", "ليسا", "ليستا", "ليسوا", "غير", "لم", "لن"}


def is_negative(toks: list[set[str]]) -> bool:
    """La phrase contient-elle une négation ? Alors la bonne réponse est celle qui est FAUSSE."""
    return any(forms & _NEGATIONS for forms in toks)


@dataclass
class GoldCheck:
    critical: list[str] = field(default_factory=list)  # réponse FAUSSE marquée juste
    rejected: list[str] = field(default_factory=list)  # bonne réponse marquée fausse
    tested: set[int] = field(default_factory=set)  # faits retrouvés dans les exercices


def check_against_gold(lesson, facts: list[dict]) -> GoldCheck:
    """Compare les exercices d'une leçon au corrigé."""
    result = GoldCheck()
    for goal in lesson.goals:
        for skill in goal.skills:
            for chunk in skill.chunks:
                if chunk.chunk_type != "practice" or not chunk.options:
                    continue
                question = tokens(chunk.content)
                question_negative = is_negative(question)
                for option in chunk.options:
                    opt = tokens(option.text)
                    # « quel groupe N'a PAS 4 éléments ? » → la bonne réponse est un fait FAUX ;
                    # une option elle-même négative (« الكتاب ليس فوق الطاولة ») inverse aussi le sens
                    negative = question_negative != is_negative(opt)
                    # une option-phrase (« السيارة على يمين الكرة ») se lit seule ; une option courte
                    # (« 3 », « مجموعة الهواتف ») se lit avec sa question (« ... عددها 4 ؟ »)
                    sequences = [(opt, 0, len(opt))]
                    alone = any(match_fact(f["words"], opt) for f in facts)
                    if len(opt) <= 4 and not alone:
                        sequences += [(question + opt, len(question), len(question) + len(opt)),
                                      (opt + question, 0, len(opt))]
                    for k, fact in enumerate(facts):
                        pages = fact.get("pages")
                        if pages and chunk.source_page is not None and chunk.source_page not in pages:
                            continue
                        if not all(any(_word_matches(c, t) for t in question) for c in fact.get("context", [])):
                            continue
                        for seq, lo, hi in sequences:
                            found = match_fact(fact["words"], seq)
                            if found and any(lo <= p < hi for p in found):  # au moins un mot dans l'option
                                result.tested.add(k)
                                where = f"p{chunk.source_page} « {chunk.content[:60]} » → « {option.text[:50]} »"
                                # dans une question négative, la bonne réponse est celle dont le fait est FAUX
                                good_answer = fact["true"] != negative
                                if option.is_correct and not good_answer:
                                    result.critical.append(where)
                                elif not option.is_correct and good_answer:
                                    result.rejected.append(where)
                                break
    return result


# --- Lancer le benchmark -----------------------------------------------------------------------
def _parse_lessons(spec: str, outline: Outline) -> list[tuple[int, int]]:
    out: list[tuple[int, int]] = []
    for part in spec.replace(" ", "").split(","):
        m = re.fullmatch(r"(\d+)\.(\d+)(?:-(\d+)\.(\d+))?", part)
        if not m:
            raise ValueError(f"--lessons invalide : « {part} ». Exemple : 1.1-1.8 ou 1.1,1.3,2.2")
        u1, l1 = int(m.group(1)), int(m.group(2))
        u2, l2 = (int(m.group(3)), int(m.group(4))) if m.group(3) else (u1, l1)
        for u in range(u1, u2 + 1):
            if not 1 <= u <= len(outline.units):
                raise ValueError(f"L'unité {u} n'existe pas dans le plan.")
            count = len(outline.units[u - 1].lessons)
            first = l1 if u == u1 else 1
            last = l2 if u == u2 else count
            for li in range(first, min(last, count) + 1):
                if (u, li) not in out:
                    out.append((u, li))
    return out


def _slug(text: str) -> str:
    return re.sub(r"[^a-zA-Z0-9._-]+", "_", text)


def run_bench(args) -> int:
    settings = get_settings()
    book = Path(args.file)
    data = book.read_bytes()
    raw = json.loads(Path(args.plan).read_text(encoding="utf-8"))
    outline = Outline.model_validate(raw.get("outline", raw))
    gold = json.loads(Path(args.gold).read_text(encoding="utf-8")) if args.gold and Path(args.gold).exists() else None
    spec = args.lessons or (",".join(gold["lessons"]) if gold else "1.1")
    lessons = _parse_lessons(spec, outline)
    common = dict(language=args.language, stage=args.stage, level=args.level, subject_name=args.subject_name,
                  retry_attempts=settings.retry_attempts, retry_wait_seconds=settings.retry_wait_seconds)

    records = []
    for spec_model in [m.strip() for m in args.models.split(",") if m.strip()]:
        provider_name, _, model = spec_model.partition(":")
        print(f"\n=== {spec_model} ===")
        try:
            ai = get_provider(settings, subject_id=0, name=provider_name, model=model or None)
        except ProviderError as exc:
            print(f"   ⏭ ignoré : {exc}")
            records.append({"model": spec_model, "status": "skipped", "reason": str(exc)})
            continue
        model = ai.model
        called, stop = False, False
        for run in range(1, args.runs + 1):
            for ui, li in lessons:
                key = f"{ui}.{li}"
                fp = lesson_fingerprint(outline, ui, li, f"{provider_name}:{model}", args.language, args.stage,
                                        args.level, args.subject_name)
                file = OUT / _slug(f"{provider_name}_{model}") / f"run{run}" / f"u{ui}-l{li}.json"
                saved = _load(file, fp)
                if saved:
                    print(f"   ↺ {key} (passage {run}) déjà fait")
                    records.append(saved | {"model": spec_model, "run": run, "lesson": key})
                    continue
                if stop:
                    records.append({"model": spec_model, "run": run, "lesson": key, "status": "quota"})
                    continue
                if called:
                    time.sleep(settings.pause_between_calls_seconds)
                called = True
                ai.reset_usage()
                started = time.monotonic()
                try:
                    result = run_lesson(data, book.name, outline, ui, li, ai, **common)
                    status, error = "ok", ""
                except QuotaExhaustedError as exc:
                    print(f"   ⛔ {key} : quota ou crédit épuisé, on passe au modèle suivant. Relancer plus tard pour finir.")
                    records.append({"model": spec_model, "run": run, "lesson": key, "status": "quota",
                                    "reason": str(exc)[:300]})
                    stop = True
                    continue
                except (TransientProviderError, PipelineError, ProviderError) as exc:
                    result, status, error = None, "failed", str(exc)[:300]
                duration = round(time.monotonic() - started, 1)
                record = {"status": status, "error": error, "duration": duration, "usage": asdict(ai.usage),
                          "result": result.model_dump() if result else None, "model_id": model}
                if status == "ok":
                    _save(file, fp, record)
                    print(f"   ✔ {key} (passage {run}) – {duration} s – {ai.usage.input_tokens} + "
                          f"{ai.usage.output_tokens} tokens")
                else:
                    print(f"   ✘ {key} (passage {run}) : {error[:150]}")
                records.append(record | {"model": spec_model, "run": run, "lesson": key})

    report = score(records, gold, args.lessons_per_book)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "resultats.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "rapport.md").write_text(markdown_report(report, args.lessons_per_book), encoding="utf-8")
    review = write_review_sheet(records, OUT)
    print("\n" + markdown_report(report, args.lessons_per_book))
    print(f"Rapport : {OUT / 'rapport.md'}  |  détails : {OUT / 'resultats.json'}")
    if review:
        print(f"Feuille de notation à l'aveugle pour l'admin : {review}")
        print(f"  puis : python3 -m app.bench note {review}")
    return 0


def _load(file: Path, fingerprint: str) -> Optional[dict]:
    if not file.exists():
        return None
    try:
        saved = json.loads(file.read_text(encoding="utf-8"))
        return saved["record"] if saved.get("fingerprint") == fingerprint else None
    except (OSError, ValueError, KeyError):
        return None


def _save(file: Path, fingerprint: str, record: dict) -> None:
    file.parent.mkdir(parents=True, exist_ok=True)
    file.write_text(json.dumps({"fingerprint": fingerprint, "record": record}, ensure_ascii=False), encoding="utf-8")


# --- Notes ----------------------------------------------------------------------------------------
def score(records: list[dict], gold: Optional[dict], lessons_per_book: int) -> dict:
    by_model: dict[str, dict] = {}
    for r in records:
        m = by_model.setdefault(r["model"], {
            "model": r["model"], "status": "ok", "lessons_done": 0, "lessons_failed": 0, "lessons_quota": 0,
            "practice": 0, "chunks": 0, "skills": 0, "critical": 0, "rejected": 0, "facts_tested": 0,
            "facts_total": 0, "rules_errors": 0, "image_needed": 0, "verify": 0, "repaired": 0,
            "duration": 0.0, "input_tokens": 0, "output_tokens": 0, "ocr_pages": 0, "cost": 0.0,
            "price_known": True, "examples_critical": [], "examples_rejected": [],
        })
        if r.get("status") == "skipped":
            m["status"] = "ignoré : " + r.get("reason", "")[:120]
            continue
        if r.get("status") == "quota":
            m["lessons_quota"] += 1
            continue
        usage = Usage(**r.get("usage", {})) if r.get("usage") else Usage()
        m["input_tokens"] += usage.input_tokens
        m["output_tokens"] += usage.output_tokens
        m["ocr_pages"] += usage.ocr_pages
        cost = cost_of(r.get("model_id", r["model"].partition(":")[2]), usage)
        if cost is None:
            m["price_known"] = False
        else:
            m["cost"] += cost
        if r.get("status") != "ok" or not r.get("result"):
            m["lessons_failed"] += 1
            continue
        try:
            res = LessonResponse.model_validate(r["result"])
        except ValidationError:
            m["lessons_failed"] += 1
            continue
        m["lessons_done"] += 1
        m["duration"] += r.get("duration", 0)
        m["repaired"] += 1 if res.meta.repaired else 0
        m["rules_errors"] += sum(1 for i in res.issues if i.level == "error")
        chunks = [c for g in res.lesson.goals for s in g.skills for c in s.chunks]
        m["skills"] += sum(len(g.skills) for g in res.lesson.goals)
        m["chunks"] += len(chunks)
        m["practice"] += sum(1 for c in chunks if c.chunk_type == "practice")
        m["image_needed"] += sum(1 for c in chunks if c.needs_image)
        m["verify"] += sum(1 for c in chunks if (c.image_note or "").strip().upper().startswith("VERIFY"))
        facts = (gold or {}).get("lessons", {}).get(r["lesson"], {}).get("facts", [])
        if facts:
            check = check_against_gold(res.lesson, facts)
            m["critical"] += len(check.critical)
            m["rejected"] += len(check.rejected)
            m["facts_tested"] += len(check.tested)
            m["facts_total"] += len(facts)
            m["examples_critical"] += [f"{r['lesson']} {e}" for e in check.critical][:10]
            m["examples_rejected"] += [f"{r['lesson']} {e}" for e in check.rejected][:10]

    for m in by_model.values():
        done = m["lessons_done"]
        calls_cost = m["cost"] if m["price_known"] else None
        attempted = done + m["lessons_failed"]
        m["cost_per_lesson"] = round(calls_cost / attempted, 4) if calls_cost is not None and attempted else None
        m["cost_per_book"] = round(m["cost_per_lesson"] * lessons_per_book, 2) if m["cost_per_lesson"] is not None else None
        m["avg_duration"] = round(m["duration"] / done, 1) if done else None
        m["critical_per_100"] = round(100 * m["critical"] / m["practice"], 1) if m["practice"] else None
        m["coverage"] = round(100 * m["facts_tested"] / m["facts_total"]) if m["facts_total"] else None
        m["cost"] = round(m["cost"], 4)
        m["price"] = price_of(m["model"].partition(":")[2])
    return {"prices_date": PRICES_DATE, "lessons_per_book": lessons_per_book, "models": list(by_model.values())}


def markdown_report(report: dict, lessons_per_book: int) -> str:
    def fmt(v, suffix=""):
        return "—" if v is None else f"{v}{suffix}"

    lines = [
        f"# Benchmark des modèles — prix du {report['prices_date']}",
        "",
        "| Modèle | Leçons OK / ratées / quota | Exercices | **Réponses fausses marquées justes** | "
        "Bonnes réponses refusées | Couverture du corrigé | Erreurs de règles | Images / VERIFY | "
        "Durée moy. (s) | Tokens entrée + sortie | Coût par leçon | **Coût par livre** |",
        "|---|---|---|---|---|---|---|---|---|---|---|---|",
    ]
    for m in report["models"]:
        if m["status"] != "ok":
            lines.append(f"| {m['model']} | {m['status']} |" + " |" * 10)
            continue
        cost_lesson = fmt(m["cost_per_lesson"] if m["cost_per_lesson"] is None else f"{m['cost_per_lesson']:.3f}", " $")
        lines.append(
            f"| {m['model']} | {m['lessons_done']} / {m['lessons_failed']} / {m['lessons_quota']} | {m['practice']} | "
            f"**{m['critical']}** ({fmt(m['critical_per_100'], ' pour 100')}) | {m['rejected']} | "
            f"{fmt(m['coverage'], ' %')} | {m['rules_errors']} | {m['image_needed']} / {m['verify']} | "
            f"{fmt(m['avg_duration'])} | {m['input_tokens']} + {m['output_tokens']} | {cost_lesson} | "
            f"**{fmt(m['cost_per_book'], ' $')}** |"
        )
    lines += [
        "",
        f"Coût par livre = coût moyen d'une leçon × {lessons_per_book} leçons (sans le plan ni les liens, "
        "qui ajoutent quelques centimes).",
        "",
        "## Exemples d'erreurs graves (à vérifier dans le livre)",
    ]
    for m in report["models"]:
        for e in m.get("examples_critical", [])[:5]:
            lines.append(f"- **{m['model']}** : {e}")
    return "\n".join(lines) + "\n"


# --- Notation à l'aveugle par l'admin ---------------------------------------------------------
DECISIONS = ("garder", "modifier", "supprimer")


def write_review_sheet(records: list[dict], out_dir: Path, seed: int = 7) -> Optional[Path]:
    """Excel : tous les chunks du 1er passage, mélangés, SANS le nom du modèle."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("(openpyxl absent : pas de feuille de notation. pip install openpyxl)")
        return None
    rows, key = [], {}
    for r in records:
        if r.get("status") != "ok" or not r.get("result") or r.get("run", 1) != 1:
            continue
        res = LessonResponse.model_validate(r["result"])
        for goal in res.lesson.goals:
            for skill in goal.skills:
                for chunk in skill.chunks:
                    options = ""
                    if chunk.options:
                        options = "\n".join(
                            f"{'✔' if o.is_correct else '✘'} {o.text}"
                            + (f"  [{o.error_category}]" if o.error_category else "") + f" — {o.feedback}"
                            for o in chunk.options
                        )
                    rows.append({"model": r["model"], "lesson": r["lesson"], "page": chunk.source_page,
                                 "type": chunk.chunk_type, "skill": skill.name, "content": chunk.content,
                                 "options": options, "image": chunk.image_note or ""})
    if not rows:
        return None
    random.Random(seed).shuffle(rows)
    rows.sort(key=lambda x: x["lesson"])  # par leçon, modèles mélangés à l'intérieur
    wb = Workbook()
    ws = wb.active
    ws.title = "Notation"
    ws.sheet_view.rightToLeft = False
    headers = ["Code", "Leçon", "Page", "Type", "Skill", "Contenu", "Options (✔ = marquée juste)",
               "Image à ajouter", "Décision : garder / modifier / supprimer", "Commentaire"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    rng = random.Random(seed + 1)
    for row in rows:
        code = f"C{rng.randrange(10**5, 10**6)}"
        while code in key:
            code = f"C{rng.randrange(10**5, 10**6)}"
        key[code] = {"model": row["model"], "lesson": row["lesson"]}
        ws.append([code, row["lesson"], row["page"], row["type"], row["skill"], row["content"], row["options"],
                   row["image"], "", ""])
    widths = [10, 7, 6, 10, 30, 60, 70, 40, 18, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row[4:8]:
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")
    choice = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    ws.add_data_validation(choice)
    choice.add(f"I2:I{ws.max_row}")
    ws.freeze_panes = "A2"
    path = out_dir / "revue_aveugle.xlsx"
    out_dir.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    # la clé (code -> modèle) est dans un fichier séparé : ne pas la donner à l'admin
    (out_dir / "cle_revue_NE_PAS_DONNER.json").write_text(json.dumps(key, ensure_ascii=False, indent=1),
                                                          encoding="utf-8")
    return path


def score_review(sheet: Path) -> dict:
    from openpyxl import load_workbook

    key = json.loads((sheet.parent / "cle_revue_NE_PAS_DONNER.json").read_text(encoding="utf-8"))
    ws = load_workbook(sheet).active
    stats: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        code, decision = row[0], (str(row[8] or "").strip().lower())
        if code not in key:
            continue
        s = stats.setdefault(key[code]["model"], {d: 0 for d in DECISIONS} | {"non noté": 0})
        s[decision if decision in DECISIONS else "non noté"] += 1
    for s in stats.values():
        noted = sum(s[d] for d in DECISIONS)
        s["gardés tels quels (%)"] = round(100 * s["garder"] / noted) if noted else None
    return stats


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Comparer les modèles d'IA sur la découpe des leçons FATIN.")
    sub = parser.add_subparsers(dest="command", required=True)
    run = sub.add_parser("run", help="Lancer le benchmark")
    run.add_argument("file", help="Le livre (PDF)")
    run.add_argument("--plan", required=True, help="Plan vérifié (JSON de l'étape 1)")
    run.add_argument("--models", default=DEFAULT_MODELS, help="fournisseur:modèle, séparés par des virgules")
    run.add_argument("--lessons", default="", help="ex. 1.1-1.8 ou 1.1,1.3 (vide = les leçons du corrigé)")
    run.add_argument("--gold", default=str(DEFAULT_GOLD), help="Corrigé de référence (JSON)")
    run.add_argument("--runs", type=int, default=1, help="Passages par modèle (2 = voir s'il est régulier)")
    run.add_argument("--lessons-per-book", type=int, default=56, help="Pour estimer le coût d'un livre")
    run.add_argument("--language", choices=["ar", "fr"], default="ar")
    run.add_argument("--stage", default="")
    run.add_argument("--level", default="")
    run.add_argument("--subject-name", default="")
    note = sub.add_parser("note", help="Compter la notation à l'aveugle remplie par l'admin")
    note.add_argument("sheet", help="resultats/bench/revue_aveugle.xlsx, rempli")
    args = parser.parse_args(argv)
    if args.command == "note":
        stats = score_review(Path(args.sheet))
        for model, s in stats.items():
            print(f"{model:40} gardés : {s['gardés tels quels (%)']} %   ({s})")
        return 0
    try:
        return run_bench(args)
    except (OSError, ValueError, ValidationError) as exc:
        print(f"Fichier ou option invalide : {exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())