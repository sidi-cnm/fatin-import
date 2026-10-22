"""Benchmark : corrigé de référence, coût, reprise, notation à l'aveugle."""
import copy
import json

import pytest

from app.bench import check_against_gold, main as bench, normalize, ordered_match, score, tokens
from app.schemas import Lesson
from tests.test_auto import EXAMPLE, book

FACTS = [
    {"pages": [12], "words": ["سياره", "يمين", "كره"], "true": False},
    {"pages": [12], "words": ["محفظه", "يمين", "كره"], "true": True},
    {"pages": [19], "words": ["4|اربع*", "هاتف|هواتف"], "true": False},
    {"pages": [16], "context": ["مجموع*"], "words": ["2|اثن*", "مثلث*"], "true": True},
]


def lesson_with(practices) -> Lesson:
    """Une leçon minimale : une skill avec les exercices donnés [(page, question, [(texte, juste)])]."""
    chunks = [{"chunk_type": "rule", "content": "r", "order_index": 1, "origin": "proposed"},
              {"chunk_type": "example", "content": "e", "order_index": 2, "origin": "proposed"}]
    for k, (page, question, options) in enumerate(practices, start=3):
        chunks.append({
            "chunk_type": "practice", "content": question, "order_index": k, "source_page": page,
            "origin": "extracted",
            "options": [{"id": "abcd"[i], "text": t, "is_correct": ok,
                         "error_category": None if ok else "concept", "feedback": "f"}
                        for i, (t, ok) in enumerate(options)],
        })
    skill = {"ref": "s1", "name": "n", "skill_type": "conceptual", "difficulty": 1, "order_index": 1,
             "origin": "proposed", "chunks": chunks}
    goals = [{"goal_type": "conceptual", "title": "g", "order_index": 1, "origin": "proposed", "skills": [skill]},
             {"goal_type": "operational", "title": "h", "order_index": 2, "origin": "proposed",
              "skills": [dict(copy.deepcopy(skill), ref="s2", skill_type="operational")]}]
    return Lesson.model_validate({"ref": "l1", "title": "t", "order_index": 1, "origin": "extracted",
                                   "goals": goals})


def test_words_without_harakat_and_without_al():
    assert normalize("الْمِحْفَظَةُ") == "المحفظه"
    toks = tokens("وَالْكُرَةُ عَلَى يَمِينِ السَّيَّارَةِ")
    assert "كره" in toks[0] and ordered_match(["كره", "يمين", "سياره"], toks) == [0, 2, 3]
    assert ordered_match(["سياره", "يمين", "كره"], toks) is None  # l'ordre compte


def test_wrong_answer_marked_correct_is_critical():
    lesson = lesson_with([
        (12, "أَلَوِّنُ الْبِطَاقَةَ الصَّحِيحَةَ", [("السَّيَّارَةُ عَلَى يَمِينِ الْكُرَةِ.", True),
                                                  ("الْمِحْفَظَةُ عَلَى يَمِينِ الْكُرَةِ.", False)]),
        (19, "أَيُّ الْمَجْمُوعَاتِ عَدَدُ عَنَاصِرِهَا 4؟", [("مَجْمُوعَةُ الْهَوَاتِفِ", True),
                                                          ("مَجْمُوعَةُ الْأَطْفَالِ", False)]),
    ])
    check = check_against_gold(lesson, FACTS)
    assert len(check.critical) == 4  # 2 skills x 2 erreurs graves
    assert len(check.rejected) == 2  # « le cartable à droite du ballon » est vrai mais marqué faux
    assert check.tested == {0, 1, 2}


def test_correct_answers_give_no_error_and_context_is_required():
    lesson = lesson_with([
        (12, "اختر", [("الْمِحْفَظَةُ عَلَى يَمِينِ الْكُرَةِ.", True), ("السَّيَّارَةُ عَلَى يَمِينِ الْكُرَةِ.", False)]),
        (16, "كم مثلثا تحيط؟ 2", [("مُثَلَّثَيْنِ", False)]),  # pas « مجموعة » dans la question : fait ignoré
        (7, "أَيْنَ؟", [("السَّيَّارَةُ عَلَى يَمِينِ الْكُرَةِ.", True)]),  # autre page : fait ignoré
    ])
    check = check_against_gold(lesson, FACTS)
    assert check.critical == [] and check.rejected == []


def test_negative_question_reverses_the_check():
    """« Quel groupe N'a PAS 4 éléments ? » → « les téléphones » est la BONNE réponse."""
    facts = FACTS + [{"pages": [19], "words": ["4|اربع*", "اطفال"], "true": True}]
    lesson = lesson_with([
        (19, "أَيُّ مَجْمُوعَةٍ لَا تَحْتَوِي عَلَى 4 عَنَاصِرَ؟", [("مَجْمُوعَةُ الْهَوَاتِفِ", True),
                                                                ("مَجْمُوعَةُ الْأَطْفَالِ", False)]),
    ])
    check = check_against_gold(lesson, facts)
    assert check.critical == [] and check.rejected == []
    wrong = lesson_with([
        (19, "أَيُّ مَجْمُوعَةٍ لَا تَحْتَوِي عَلَى 4 عَنَاصِرَ؟", [("مَجْمُوعَةُ الْأَطْفَالِ", True)]),
    ])
    assert len(check_against_gold(wrong, facts).critical) == 2  # les enfants SONT 4 : erreur grave


def test_two_word_fact_matches_in_any_order_but_three_word_fact_keeps_order():
    facts = [{"pages": [16], "context": ["مجموع*"], "words": ["2|اثن*", "تفاح*"], "true": False},
             {"pages": [12], "words": ["سياره", "يمين", "كره"], "true": False}]
    lesson = lesson_with([
        (16, "أُلَوِّنُ الْمَجْمُوعَةَ الَّتِي يُسَاوِي عَنَاصِرُهَا اثْنَيْنِ",
         [("تلوين المجموعة التي تحتوي على تفاحتين اثنتين", True)]),  # « pommes deux » : ordre inversé
        (12, "اختر", [("الْكُرَةُ عَلَى يَمِينِ السَّيَّارَةِ.", True)]),  # autre sens : ne doit PAS compter
    ])
    check = check_against_gold(lesson, facts)
    assert len(check.critical) == 2 and all("p16" in e for e in check.critical)


def test_score_computes_cost_per_book():
    lesson = lesson_with([(12, "اختر", [("الْمِحْفَظَةُ عَلَى يَمِينِ الْكُرَةِ.", True)])])
    result = {"unit_index": 1, "lesson_index": 3, "lesson": lesson.model_dump(), "issues": [],
              "meta": {"provider": "gemini", "model": "gemini-3.8-flash", "duration_seconds": 10, "repaired": False}}
    records = [{"model": "gemini:gemini-3.8-flash", "model_id": "gemini-3.8-flash", "run": 1, "lesson": "1.3",
                "status": "ok", "duration": 10, "result": result,
                "usage": {"calls": 1, "input_tokens": 5000, "output_tokens": 8000, "ocr_pages": 0}}]
    gold = {"lessons": {"1.3": {"facts": FACTS}}}
    m = score(records, gold, 56)["models"][0]
    per_lesson = 5000 / 1e6 * 0.75 + 8000 / 1e6 * 3.75
    assert m["cost_per_lesson"] == pytest.approx(per_lesson, abs=1e-4)
    assert m["cost_per_book"] == pytest.approx(round(per_lesson, 4) * 56, abs=0.01)
    assert m["critical"] == 0 and m["lessons_done"] == 1


@pytest.fixture
def bench_dir(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("PAUSE_BETWEEN_CALLS_SECONDS", "0")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    (tmp_path / "livre.pdf").write_bytes(book(20))
    (tmp_path / "plan.json").write_text(EXAMPLE.model_dump_json(), encoding="utf-8")
    gold = {"lessons": {"1.1": {"facts": FACTS}, "1.2": {"facts": FACTS}}}
    (tmp_path / "gold.json").write_text(json.dumps(gold), encoding="utf-8")
    return tmp_path


def test_bench_run_resume_and_blind_review(bench_dir, capsys):
    args = ["run", "livre.pdf", "--plan", "plan.json", "--gold", "gold.json",
            "--models", "fake:sample,openai:gpt-6.1-sol"]
    assert bench(args) == 0
    out = capsys.readouterr().out
    assert "ignoré" in out and "OPENAI_API_KEY" in out  # pas de clé : modèle ignoré, les autres continuent
    report = json.loads((bench_dir / "resultats" / "bench" / "resultats.json").read_text(encoding="utf-8"))
    fake = report["models"][0]
    assert fake["lessons_done"] == 2 and fake["practice"] > 0
    assert (bench_dir / "resultats" / "bench" / "rapport.md").exists()

    assert bench(args) == 0  # 2e fois : reprise, rien n'est redemandé
    assert capsys.readouterr().out.count("déjà fait") == 2

    from openpyxl import load_workbook

    sheet = bench_dir / "resultats" / "bench" / "revue_aveugle.xlsx"
    wb = load_workbook(sheet)
    ws = wb.active
    assert "fake" not in json.dumps([[c.value for c in r] for r in ws.iter_rows()])  # à l'aveugle
    for i, row in enumerate(ws.iter_rows(min_row=2), start=0):
        row[8].value = "garder" if i % 2 == 0 else "modifier"
    wb.save(sheet)
    assert bench(["note", str(sheet)]) == 0
    assert "gardés" in capsys.readouterr().out